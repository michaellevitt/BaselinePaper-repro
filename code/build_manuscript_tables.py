#!/usr/bin/env python3
"""Every table in the manuscript, with the manuscript's own numbering, in one workbook:
docs/PaperA_manuscript_tables.xlsx

  Table 1  Per-population pre-pandemic trend patterns (TTa slope and slope-of-slopes with
           p-values, shrinkage weight, STTa slope and slope-of-slopes)
  Table 2  Combined excess deaths + pooled P-score, by period x age group x trend modeling x model
  Table 3  Excess mortality, more- vs less-vulnerable populations, by period x model
  Table S1 TTa excess with age-band versus country-level trends, 10 largest countries
  Table S2 Excess deaths by single year and model, all ages, pooled
  Table S3 Per-population excess P-score, all ages (3 periods x 4 models + mean of 4)
  Table S4 ... 65 and over
  Table S5 ... under 65
  Table S6 Effect of age adjustment (age-resolution) on the 2020-2025 total
  Table S7 Vulnerability indicators and their sources
  Table S8 England and Wales versus the whole United Kingdom

Reads the engine outputs in output/ plus data/vuln_covariates_38_v3.csv.
Run run_all.sh, which produces every input in the required order (the two engines, the
age-band trends, the Table S8 script, build_option1_age_resolution.py and
build_table_S6_agespecific.py)."""
import os, csv, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
ROOT=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTD=os.path.join(ROOT,"output"); DATA=os.path.join(ROOT,"data"); DOCS=os.path.join(ROOT,"docs")
HDR=PatternFill("solid",fgColor="1F4E79"); WHT=Font(color="FFFFFF",bold=True)
SUM=PatternFill("solid",fgColor="FFF2CC")
thin=Side(style="thin",color="BFBFBF"); BORD=Border(left=thin,right=thin,top=thin,bottom=thin)
CEN=Alignment(horizontal="center",wrap_text=True)
MODELS=["Fa","TTa","STTa","STTa+"]
GROUP_LABEL={"All":"All ages","GE65":"65+ years","LT65":"<65 years"}
PERIODS=[("2020-2025","2020-2025"),("2020-2023","2020-2023"),("2024-2025","2024-2025")]

def _f(x):
    try: return float(x)
    except (TypeError,ValueError): return None
wb=Workbook(); wb.remove(wb.active)
def sheet(name,header,rows,widths=None,sumrows=0):
    ws=wb.create_sheet(name); ws.append(header)
    for j in range(1,len(header)+1):
        c=ws.cell(1,j); c.fill=HDR; c.font=WHT; c.alignment=CEN; c.border=BORD
    for r in rows:
        ws.append(r); rn=ws.max_row
        for j in range(1,len(header)+1):
            c=ws.cell(rn,j); c.border=BORD
            if j>1: c.alignment=Alignment(horizontal="center")
        if sumrows and rn>len(rows)+1-sumrows:
            for j in range(1,len(header)+1): ws.cell(rn,j).fill=SUM; ws.cell(rn,j).font=Font(bold=True)
    for i,wd in enumerate(widths or [],1): ws.column_dimensions[get_column_letter(i)].width=wd
    ws.freeze_panes="A2"; return ws

# ---------- load engine outputs ----------
par=list(csv.DictReader(open(os.path.join(OUTD,"option1_country_params.csv"))))
per=list(csv.DictReader(open(os.path.join(OUTD,"model_option1_periods.csv"))))
pool=list(csv.DictReader(open(os.path.join(OUTD,"pooled_option1_by_period.csv"))))
cov=list(csv.DictReader(open(os.path.join(DATA,"vuln_covariates_38_v3.csv"))))
NAME={r["country"]:r["name"] for r in par}
GRP={r["code"]:r["group"] for r in cov}
E={(r["country"],r["group"],r["model"],r["period"]):(float(r["O"]),float(r["E"])) for r in per}
POOL={(r["period"],r["group"],r["model"]):r for r in pool}
def P(c,g,m,p):
    O,Ee=E[(c,g,m,p)]; return 100*(O-Ee)/Ee

# ---------- Table 1 ----------
# Exactly the manuscript's layout.  p-values are two-sided t tests with n-2 degrees of freedom,
# from the t-based 95% CIs in slope_of_slopes_CI.csv.  The STTa columns are
# w*own + (1-w)*aggregate in %/yr, with the aggregate the full-precision population-weighted
# mean over the 37 fitted populations (the 3-decimal value in slope_of_slopes_CI.csv shifts the
# last printed digit for some countries).  Native float formatting, as in the manuscript.
import scipy.stats as _ss
CI={r["country"]:r for r in csv.DictReader(open(os.path.join(OUTD,"slope_of_slopes_CI.csv")))
    if r["country"] and not r["country"].startswith(("WEIGHTED","UNWEIGHTED","SD"))}
T1NAME={"The United States of America":"United States","Bulgaria":"Bulgaria*"}
def _pv(est,lo,hi,n):
    dof=n-2; t975=_ss.t.ppf(0.975,dof); se=(hi-lo)/(2*t975)
    p=2*_ss.t.sf(abs(est/se),dof)
    return "<0.001" if p<0.001 else f"{p:.3f}"
fit=[r for r in par if _f(r["islope2019"]) is not None]
_tot=sum(_f(r["pop2019"]) for r in fit)
AGG_G=sum(_f(r["pop2019"])*_f(r["islope2019"]) for r in fit)/_tot
AGG_Q=sum(_f(r["pop2019"])*_f(r["sos"]) for r in fit)/_tot
rows=[]; sg_sum=sq_sum=0.0
for r in par:
    name=T1NAME.get(r["name"],r["name"]); wv=_f(r["w"]) or 0.0
    g=_f(r["islope2019"]); q=_f(r["sos"])
    if g is None:
        rows.append([name,"\u2014","\u2014","\u2014","\u2014",f"{wv:.2f}",f"{AGG_G:+.2f}",f"{AGG_Q:+.3f}"])
        continue
    c=CI[r["country"]]; n=int(c["n_pre"])
    sg=wv*g+(1-wv)*AGG_G; sq=wv*q+(1-wv)*AGG_Q
    sg_sum+=_f(r["pop2019"])*sg; sq_sum+=_f(r["pop2019"])*sq
    rows.append([name,f"{g:+.2f}",_pv(float(c["u_islope2019"]),float(c["u_islope2019_lo"]),float(c["u_islope2019_hi"]),n),
                 f"{q:+.3f}",_pv(float(c["u_sos"]),float(c["u_sos_lo"]),float(c["u_sos_hi"]),n),
                 f"{wv:.2f}",f"{sg:+.2f}",f"{sq:+.3f}"])
rows.sort(key=lambda x: x[0])
rows.append([f"Weighted mean (n={len(fit)})",f"{AGG_G:+.2f}","",f"{AGG_Q:+.3f}","","",
             f"{sg_sum/_tot:+.2f}",f"{sq_sum/_tot:+.3f}"])
sheet("Table 1",["Population","Slope in 2019 for TTa model (%/yr)","p-value",
                 "Slope-of-slopes for TTa model (%/yr\u00b2)","p-value","Shrinkage weight",
                 "Slope in 2019 for STTa model (%/yr)","Slope-of-slopes for STTa model (%/yr\u00b2)"],
      rows,[22,14,9,16,9,11,14,16],sumrows=1)

# ---------- Table 2 ----------
# Each period and age group twice: age-band trend modeling (the headline engine) and
# country-level trend modeling (agespecific_vs_common_v1.csv, whose country-level arm comes from
# the country-level engine).  Fa has no trend, so its two rows are identical.
from decimal import Decimal as _D, ROUND_HALF_UP as _HU
def _x(e,p):
    m=(_D(e)/_D(1000000)).quantize(_D("0.01"),rounding=_HU)
    q=_D(str(p)).quantize(_D("0.1"),rounding=_HU)
    return f"{m:+.2f}M ({q:+.1f}%)"
VSC={(r["period"],r["group"],r["model_common"]):r
     for r in csv.DictReader(open(os.path.join(OUTD,"agespecific_vs_common_v1.csv")))}
if sum(abs(float(r["delta_P_pp"]))<1e-9 for r in VSC.values())>2:
    raise SystemExit("agespecific_vs_common_v1.csv compares the age-band results with themselves: "
                     "the country-level engine did not run before model_agespecific_trends_v1.py")
EN="\u2013"
t2=[]
for pk,plabel in PERIODS:
    for g in ("All","GE65","LT65"):
        for mode in ("Age-band","Country-level"):
            row=[plabel.replace("-",EN),GROUP_LABEL[g],mode]
            fa=POOL[(pk,g,"Fa")]; row.append(_x(fa["excess"],fa["pooled_Ppct"]))
            for m in MODELS[1:]:
                if mode=="Age-band":
                    r=POOL[(pk,g,m)]; row.append(_x(r["excess"],r["pooled_Ppct"]))
                else:
                    r=VSC[(pk,g,m)]; row.append(_x(r["excess_common"],r["P_common"]))
            t2.append(row)
sheet("Table 2",["Period","Age group","Trend modeling"]+MODELS,t2,[13,12,15]+[18]*4)

# ---------- Table 3 ----------
MORE=[c for c in NAME if GRP.get(c)=="more"]; LESS=[c for c in NAME if GRP.get(c)=="less"]
def pooledP(codes,g,m,p):
    O=sum(E[(c,g,m,p)][0] for c in codes); Ee=sum(E[(c,g,m,p)][1] for c in codes)
    return 100*(O-Ee)/Ee
t3=[]
for pk,plabel in PERIODS:
    for m in MODELS:
        a=pooledP(MORE,"All",m,pk); b=pooledP(LESS,"All",m,pk)
        t3.append([plabel.replace("-","\u2013"),m,f"{a:+.2f}",f"{b:+.2f}",f"{a-b:+.2f}"])
sheet("Table 3",["Period","Model",f"More vulnerable (n={len(MORE)})",f"Less vulnerable (n={len(LESS)})","Difference (pp)"],
      t3,[13,8,20,20,16])

# ---------- Tables S3 / S4 / S5 ----------
def si_table(group,name):
    order=sorted(NAME,key=lambda c: np.mean([P(c,group,m,"2020-2025") for m in MODELS]),reverse=True)
    rows=[]
    for c in order:
        row=[NAME[c]]
        for pk,_ in PERIODS:
            vals=[P(c,group,m,pk) for m in MODELS]
            row+= [round(v,1) for v in vals]+[round(float(np.mean(vals)),1)]
        rows.append(row)
    tot=["Pooled (38 populations)"]
    for pk,_ in PERIODS:
        vals=[]
        for m in MODELS:
            r=POOL[(pk,group,m)]; vals.append(float(r["pooled_Ppct"]))
        tot += [round(v,1) for v in vals]+[round(float(np.mean(vals)),1)]
    rows.append(tot)
    hdr=["Population"]+[f"{plabel} {m}" for pk,plabel in PERIODS for m in MODELS+["mean of 4"]]
    sheet(name,hdr,rows,[22]+[11]*15,sumrows=1)
si_table("All","Table S3"); si_table("GE65","Table S4"); si_table("LT65","Table S5")

# ---------- Table S2 (single years) ----------
YEARS=[str(y) for y in range(2020,2026)]
t4=[]
for m in MODELS:
    row=[m]
    for y in YEARS:
        r=POOL[(y,"All",m)]
        row.append(f"{float(r['excess'])/1e3:+.0f} ({float(r['pooled_Ppct']):+.1f})")
    t4.append(row)
sheet("Table S2",["Model"]+YEARS,t4,[10]+[15]*6)

# ---------- Table S6 (age resolution) ----------
ar=list(csv.reader(open(os.path.join(OUTD,"age_resolution_option1.csv"))))
hdr=ar[0]; arm={r[0]:r for r in ar[1:]}
sheet("Table S6",[hdr[0]+" (millions)"]+hdr[1:],
      [[m]+[f"{float(arm[m][i])/1e6:.2f}" for i in range(1,len(hdr))] for m in MODELS if m in arm],[22,14,16,15,17])

# ---------- Table S1 (10 largest countries) ----------
s6p=os.path.join(OUTD,"table_S6_agespecific_vs_allage.csv")
if os.path.exists(s6p):
    s6=list(csv.DictReader(open(s6p)))
    rows=[[r["name"] if r["code"]!="TOTAL" else "Total (10 countries)"]+
          [int(r[f"excess_{g}_{k}"]) for g in ("All","<65","65+") for k in ("agespecific","allage")] for r in s6]
    sheet("Table S1",["Population","All: age-specific","All: all-age","<65: age-specific","<65: all-age","65+: age-specific","65+: all-age"],
          rows,[22]+[15]*6,sumrows=1)

# ---------- Table S7 ----------
sheet("Table S7",["Code","Population","GDP per capita 2021 (US$)","Gini","Poverty (%)","Group","Trigger","Source"],
      [[r["code"],r["name"],r["gdp_pc_2021_usd"],r["gini"],r["poverty_pct"],r["group"],r.get("trigger",""),r.get("source","")]
       for r in sorted(cov,key=lambda r:r["name"])],[7,20,16,8,10,10,14,40])

# ---------- Table S8 ----------
s8p=os.path.join(OUTD,"uk_vs_ew_full_family_2020_2025_v2_6dp.csv")
MINUS="\u2212"
def _s8(x):
    v=_D(x).quantize(_D("0.01"),rounding=_HU)
    return "+0.00" if v==0 else ("+" if v>0 else MINUS)+f"{abs(v):.2f}"
if os.path.exists(s8p):
    S8={(r["period"],r["model"]):r for r in csv.DictReader(open(s8p)) if r["group"]=="All"}
    rows=[]
    for pk,plabel in PERIODS:
        for k,m in enumerate(MODELS):
            r=S8[(pk,m)]
            rows.append([plabel.replace("-",EN) if k==0 else "",m,_s8(r["ew_pooled_Ppct"]),
                         _s8(r["uk_pooled_Ppct"]),_s8(r["delta_Ppct"])])
    sheet("Table S8",["Period","Model","E&W","UK","\u0394"],rows,[13,8,10,10,10])

# sheets in the manuscript's order
_order=["Table 1","Table 2","Table 3"]+[f"Table S{i}" for i in range(1,9)]
wb._sheets.sort(key=lambda ws: _order.index(ws.title) if ws.title in _order else 99)

os.makedirs(DOCS,exist_ok=True)
out=os.path.join(DOCS,"PaperA_manuscript_tables.xlsx"); wb.save(out)
print("wrote",out)
print("sheets:",[ws.title for ws in wb.worksheets])
print("\nTable 2, 2020-2025 all ages:", t2[0][2:])
print("Table 3, 2020-2025 Fa      :", t3[0])
print("Table S2, Fa               :", t4[0][1:])
