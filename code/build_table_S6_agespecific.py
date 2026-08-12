#!/usr/bin/env python3
"""
Table S6 + Figure-S2 data — TTa excess deaths under AGE-BAND-SPECIFIC vs ALL-AGE (common) slopes and
slopes-of-slopes, for the 10 countries with the most deaths, 2020-2025, for All / <65 / 65+.

For each country and 5-year band a, the TTa baseline is
    m(a,y) = exp(alpha_a) * clip( exp(beta*t + gamma*t^2), 0.4, 2.5 ),   t = y - 2019
with alpha_a = the band's fitted 2019 log-level (deg-2 polyfit of its log rate over the reliable
pre-pandemic years — identical to the main engine).  The two variants differ only in (beta, gamma):
  * ALL-AGE  (common): the country's single two-step trend (from output/slope_of_slopes_CI.csv) — this
    column reproduces the main-engine TTa excess exactly.
  * AGE-BAND-SPECIFIC : each band's OWN two-step trend (beta_a = g19_a/100, gamma_a = q_a/200), where
    g19_a and q_a come from the per-band moving-slope OLS (same construction as Figure S2 /
    fig_perband_trend_heterogeneity.py).

Reads : data/master_5x1_DPM_90plus.csv, output/slope_of_slopes_CI.csv
Writes: output/table_S6_agespecific_vs_allage.csv
        docs/Table_S6_agespecific_vs_allage.xlsx
Prints: the paragraph-57 "vs" numbers, the 10-country totals, and every >50k divergence (comment 191).
"""
import os, sys, csv, numpy as np
HERE=os.path.dirname(os.path.abspath(__file__)); ROOT=os.path.dirname(HERE)
OUTD=os.path.join(ROOT,"output"); DOCS=os.path.join(ROOT,"docs")
sys.path.insert(0, HERE)
from excess_anchor_window import load, cell_for

FINE=["0","1-4","5-9","10-14","15-19","20-24","25-29","30-34","35-39","40-44",
      "45-49","50-54","55-59","60-64","65-69","70-74","75-79","80-84","85-89","90+"]
def astart(a): return 0 if a=="0" else (1 if a=="1-4" else int(a.split("-")[0].replace("+","")))
GE65=[a for a in FINE if astart(a)>=65]; LT65=[a for a in FINE if astart(a)<65]
GROUPS={"All":FINE,"<65":LT65,"65+":GE65}
CLIP=(0.4,2.5); PYEARS=[2020,2021,2022,2023,2024,2025]
RELIABLE={"BGR":2015,"CHL":2011,"HRV":2007,"EST":2005,"HUN":2006,"LVA":2007,"LTU":2006,"POL":2008,"SVK":2006}
# the 10 countries with most deaths, in John's manuscript order
TEN=[("USA","USA"),("JPN","Japan"),("DEUTNP","Germany"),("ITA","Italy"),("GBRTENW","England/Wales"),
     ("FRATNP","France"),("ESP","Spain"),("POL","Poland"),("KOR","Korea"),("CAN","Canada")]

data,name=load()
def DP(l,y,a): return cell_for(data,l,y,a,"T")
def baseyrs(l): return list(range(max(2003,RELIABLE.get(l,2003)),2020))
def has_year(l,y): return all(DP(l,y,a) is not None for a in FINE)

# common (all-age) two-step trend per country
COMMON={}
for r in csv.DictReader(open(os.path.join(OUTD,"slope_of_slopes_CI.csv"))):
    if r.get("country") and r.get("u_islope2019","") not in ("",None) and r.get("u_sos","") not in ("",None):
        COMMON[r["country"]]=(float(r["u_islope2019"])/100.0, float(r["u_sos"])/200.0)

def alpha(l,a):                                  # band's fitted 2019 log-level (matches engine)
    ys=baseyrs(l); t=np.array(ys,float)-2019.0
    r=np.array([DP(l,y,a)[0]/DP(l,y,a)[1] for y in ys])
    return np.polyfit(t,np.log(np.maximum(r,1e-9)),2)[2]

def mslope(l,a,c):                               # centred 5-yr log-slope (%/yr) at window centre c
    pts=[(y,DP(l,y,a)) for y in range(c-2,c+3)]; pts=[(y,cc) for y,cc in pts if cc and cc[1]>0 and cc[0]>0]
    if len(pts)<3: return None
    x=np.array([p[0] for p in pts],float); yv=np.log(np.array([p[1][0]/p[1][1] for p in pts])); xc=x-x.mean()
    return 100.0*(xc*(yv-yv.mean())).sum()/(xc**2).sum()

def perband_trend(l,a):                          # each band's own two-step (beta_a, gamma_a)
    rsd=RELIABLE.get(l,2003)
    pre=[(c,mslope(l,a,c)) for c in range(rsd+2,2018)]; pre=[(c,v) for c,v in pre if v is not None]
    if len(pre)<3: return None
    x=np.array([c for c,_ in pre],float); y=np.array([v for _,v in pre]); xb=x.mean(); Sxx=((x-xb)**2).sum()
    b=((x-xb)*(y-y.mean())).sum()/Sxx; a0=y.mean()-b*xb
    g19=a0+b*2019.0                              # fitted slope@2019 (%/yr);  b = slope-of-slopes (%/yr^2)
    return g19/100.0, b/200.0

def rate(l,a,bg):                                # TTa baseline factory for a given (beta,gamma)
    al=alpha(l,a); beta,gamma=bg
    def f(y):
        t=y-2019.0
        return np.exp(al)*min(max(np.exp(beta*t+gamma*t*t),CLIP[0]),CLIP[1])
    return f

def excess(l,trend):                             # trend: 'common' or 'agespec' -> dict group->(O,E)
    yrs=[y for y in PYEARS if has_year(l,y)]
    out={}
    for gname,bands in GROUPS.items():
        O=E=0.0
        for a in bands:
            bg=COMMON[l] if trend=="common" else perband_trend(l,a)
            if bg is None: bg=COMMON[l]           # fall back to common if a band lacks its own trend
            f=rate(l,a,bg)
            for y in yrs:
                O+=DP(l,y,a)[0]; E+=f(y)*DP(l,y,a)[1]
        out[gname]=(O,E)
    return out

rows=[]; tot={("agespec",g):[0.0,0.0] for g in GROUPS}; tot.update({("common",g):[0.0,0.0] for g in GROUPS})
big=[]
print(f"{'country':13} {'group':4} {'age-specific':>14} {'all-age':>12} {'diff':>10}")
for code,disp in TEN:
    AS=excess(code,"agespec"); CO=excess(code,"common")
    row={"code":code,"name":disp}
    for g in GROUPS:
        xa=AS[g][0]-AS[g][1]; xc=CO[g][0]-CO[g][1]
        row[f"as_{g}"]=xa; row[f"all_{g}"]=xc; row[f"diff_{g}"]=xa-xc
        tot[("agespec",g)][0]+=xa; tot[("common",g)][0]+=xc
        if abs(xa-xc)>=50000: big.append((disp,g,xa,xc,xa-xc))
        print(f"{disp:13} {g:4} {xa:>14,.0f} {xc:>12,.0f} {xa-xc:>10,.0f}")
    rows.append(row)

# ---- write CSV ----
os.makedirs(OUTD,exist_ok=True)
with open(os.path.join(OUTD,"table_S6_agespecific_vs_allage.csv"),"w",newline="") as f:
    w=csv.writer(f)
    w.writerow(["code","name",
        "excess_All_agespecific","excess_All_allage","diff_All",
        "excess_<65_agespecific","excess_<65_allage","diff_<65",
        "excess_65+_agespecific","excess_65+_allage","diff_65+"])
    for r in rows:
        w.writerow([r["code"],r["name"]]+[round(r[f"{k}_{g}"]) for g in GROUPS for k in ("as","all","diff")])
    w.writerow(["TOTAL","10 countries"]+[round(v[0]) for g in GROUPS for v in
               (tot[("agespec",g)],tot[("common",g)],[tot[("agespec",g)][0]-tot[("common",g)][0]])])

# ---- xlsx ----
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb=openpyxl.Workbook(); ws=wb.active; ws.title="Table S6"
    thin=Side(style="thin",color="BFBFBF"); B=Border(*(thin,)*4)
    HDR=PatternFill("solid",fgColor="1F4E79"); WH=Font(color="FFFFFF",bold=True)
    ws.append(["","","All ages","","<65","","65+",""])
    ws.append(["Code","Population","Age-specific","All-age","Age-specific","All-age","Age-specific","All-age"])
    for r in rows:
        ws.append([r["code"],r["name"]]+[round(r[f"{k}_{g}"]) for g in GROUPS for k in ("as","all")])
    ws.append(["TOTAL","10 countries"]+[round(tot[(t,g)][0]) for g in GROUPS for t in ("agespec","common")])
    for j in range(1,9):
        for i in (1,2): c=ws.cell(i,j); c.fill=HDR; c.font=WH; c.alignment=Alignment(horizontal="center",wrap_text=True)
    for row in ws.iter_rows():
        for c in row:
            c.border=B
            if c.column>2 and isinstance(c.value,int): c.number_format="#,##0"; c.alignment=Alignment(horizontal="right")
    ws.column_dimensions["B"].width=16
    os.makedirs(DOCS,exist_ok=True); wb.save(os.path.join(DOCS,"Table_S6_agespecific_vs_allage.xlsx"))
except Exception as e:
    print("xlsx skipped:",e)

# ---- paragraph-57 fill + comment-191 divergence list ----
def fm(x): return f"{x/1e3:,.0f}k" if abs(x)>=1000 else f"{x:,.0f}"
print("\n=== ¶57 fill (2020-2025 All-ages, age-specific vs all-age) ===")
for r in rows:
    print(f"  {r['name']:14}: {fm(r['as_All'])} vs {fm(r['all_All'])}")
print(f"  {'TOTAL(10)':14}: {fm(tot[('agespec','All')][0])} vs {fm(tot[('common','All')][0])}")
print("\n=== comment-191: divergences >= 50k (country, stratum, age-spec, all-age, diff) ===")
for d in sorted(big,key=lambda x:-abs(x[4])):
    print(f"  {d[0]:14} {d[1]:4}: {d[2]:>12,.0f} vs {d[3]:>12,.0f}  (diff {d[4]:+,.0f})")
print("\nwrote output/table_S6_agespecific_vs_allage.csv, docs/Table_S6_agespecific_vs_allage.xlsx")
