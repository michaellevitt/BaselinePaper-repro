#!/usr/bin/env python3
"""All Paper A tables (two-step trend, slope-anchored STTa+) as one workbook: docs/PaperA_tables_twostep_v1.xlsx.
T1 methods comparison · T2 slope table · T3 pooled excess · S1 age resolution · S2 excess by location ·
S3 excess by period · S4 vulnerability covariates · S5 wealth correlations."""
import os, csv, json, numpy as np
from scipy.stats import spearmanr, pearsonr
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
ROOT=os.path.dirname(os.path.dirname(os.path.abspath(__file__))); OUTD=os.path.join(ROOT,"output"); DATA=os.path.join(ROOT,"data"); DOCS=os.path.join(ROOT,"docs")
HDR=PatternFill("solid",fgColor="1F4E79"); WHT=Font(color="FFFFFF",bold=True); SUM=PatternFill("solid",fgColor="FFF2CC")
thin=Side(style="thin",color="BFBFBF"); BORD=Border(left=thin,right=thin,top=thin,bottom=thin); CEN=Alignment(horizontal="center",wrap_text=True)
def _f(x):
    try: return float(x)
    except (TypeError,ValueError): return None
SPECS=[]   # (short, caption, header, rows) for docx emission
def sheet(wb,name,header,rows,widths=None,sumrows=0,caption=None):
    SPECS.append((name,caption or name,list(header),[list(r) for r in rows]))
    ws=wb.create_sheet(name); ws.append(header)
    for j in range(1,len(header)+1): c=ws.cell(1,j); c.fill=HDR; c.font=WHT; c.alignment=CEN; c.border=BORD
    for r in rows:
        ws.append(r); rn=ws.max_row
        for j in range(1,len(header)+1):
            c=ws.cell(rn,j); c.border=BORD
            if j>1: c.alignment=Alignment(horizontal="center")
        if sumrows and rn>len(rows)-sumrows+1:
            for j in range(1,len(header)+1): ws.cell(rn,j).fill=SUM; ws.cell(rn,j).font=Font(bold=True)
    for i,w in enumerate(widths or [],1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes="A2"; return ws

wb=Workbook(); wb.remove(wb.active)
MODELS=["Fa","TTa","STTa","STTa+"]

# ---- T1 methods comparison ----
st=json.load(open(os.path.join(DATA,"methods_comparison.json")))["studies"]   # curated methods-landscape (Table 1)
sheet(wb,"T1 methods comparison",["Ref","Study","Model family","Baseline construction","Age handling","Reference period","Scope"],
      [[s["ref"],s["study"],s["family"],s["baseline"],s["age"],s["refperiod"],s["scope"]] for s in st],
      [5,22,16,34,30,20,24])

# ---- T2 slope table ----
par=list(csv.DictReader(open(os.path.join(OUTD,"option1_country_params.csv"))))
def ci(v,lo,hi): return f"{v:+.2f} ({lo:+.2f}, {hi:+.2f})" if None not in (v,lo,hi) else (f"{v:+.2f}" if v is not None else "")
rows=[]; IS=[];SOS=[];POP=[];Wl=[]
for r in par:
    is19=_f(r["islope2019"]); sos=_f(r["sos"]); pop=_f(r["pop2019"]); w=_f(r["w"])
    rows.append([r["name"],round(pop/1e6,2) if pop else "",round(w,3) if w is not None else "",
        ci(is19,_f(r["islope2019_lo"]),_f(r["islope2019_hi"])),ci(sos,_f(r["sos_lo"]),_f(r["sos_hi"])),
        f"{100*_f(r['beta_shrunk']):+.2f}" if r["beta_shrunk"] else "",f"{200*_f(r['gamma_shrunk']):+.3f}" if r["gamma_shrunk"] else "",
        f"{100*_f(r['beta_anch_shrunk']):+.2f}" if r["beta_anch_shrunk"] else "",f"{200*_f(r['gamma_anch_shrunk']):+.3f}" if r["gamma_anch_shrunk"] else ""])
    if is19 is not None: IS.append(is19);SOS.append(sos);POP.append(pop);Wl.append(w)
IS=np.array(IS);SOS=np.array(SOS);POP=np.array(POP)
rows.append(["Pop-weighted mean","","",f"{np.average(IS,weights=POP):+.2f}",f"{np.average(SOS,weights=POP):+.3f}","","","",""])
rows.append(["Unweighted mean","",f"{np.mean(Wl):.3f}",f"{IS.mean():+.2f}",f"{SOS.mean():+.3f}","","","",""])
sheet(wb,"T2 slope table",["Population","Pop (M)","w","islope@2019 [95% CI] (%/yr)","slope-of-slopes [95% CI] (%/yr²)","STTa islope","STTa sos","STTa+ islope","STTa+ sos"],
      rows,[19,7,7,25,26,10,10,10,10],sumrows=2)

# ---- pooled excess ----
pool={(r["model"],r["period"]):r for r in csv.DictReader(open(os.path.join(OUTD,"pooled_option1_by_period.csv"))) if r["group"]=="All"}
PER=[("2020-2023","2020–2023 (pandemic)"),("2024-2025","2024–2025 (post)"),("2020-2025","2020–2025 (all)")]
# ---- T3 pooled excess by baseline x period ----
t3=[]
for m in MODELS:
    row=[m]
    for pk,_ in PER:
        r=pool[(m,pk)]; row+=[round(float(r["excess"])/1e6,2),f"{float(r['pooled_Ppct']):+.1f}%"]
    t3.append(row)
sheet(wb,"T3 pooled excess",["Baseline"]+[x for _,lab in PER for x in [f"{lab}\nM",f"{lab}\nP%"]],t3,[10]+[13]*6)

# ---- S1 age resolution ----
ar=list(csv.reader(open(os.path.join(OUTD,"age_resolution_option1.csv"))))
hdr=ar[0]; arm={r[0]:r for r in ar[1:]}
sheet(wb,"S1 age resolution",[hdr[0]+" (millions)"]+hdr[1:],
      [[m]+[f"{float(arm[m][i])/1e6:.2f}" for i in range(1,len(hdr))] for m in MODELS if m in arm],[20,13,15,14,16])

# ---- per-country excess ----
grp={r["code"]:r["group"] for r in csv.DictReader(open(os.path.join(DATA,"vuln_covariates_38_v3.csv")))}
nm={r["country"]:r["name"] for r in par}
E={}
for r in csv.DictReader(open(os.path.join(OUTD,"model_option1_periods.csv"))):
    if r["group"]=="All": E[(r["country"],r["model"],r["period"])]=(float(r["O"]),float(r["E"]))
def P(c,m,per): O,Ee=E[(c,m,per)]; return 100*(O-Ee)/Ee
order=sorted(nm,key=lambda c:P(c,"STTa+","2020-2025"),reverse=True)
# ---- S2 excess by location 2020-2025 ----
sheet(wb,"S2 excess by location",["Population","Group"]+[f"{m}\nP%" for m in MODELS],
      [[nm[c],grp.get(c,"")]+[round(P(c,m,"2020-2025"),1) for m in MODELS] for c in order],[19,11]+[9]*4)
# ---- S3 excess by period ----
sheet(wb,"S3 excess by period",["Population","STTa 2020–23","STTa 2024–25","STTa+ 2020–23","STTa+ 2024–25"],
      [[nm[c],round(P(c,"STTa","2020-2023"),1),round(P(c,"STTa","2024-2025"),1),round(P(c,"STTa+","2020-2023"),1),round(P(c,"STTa+","2024-2025"),1)] for c in order],
      [19,12,12,12,12])

# ---- S4 vulnerability covariates ----
cov=list(csv.DictReader(open(os.path.join(DATA,"vuln_covariates_38_v3.csv"))))
sheet(wb,"S4 vulnerability",["Code","Population","GDP pc 2021 (USD)","Gini","Poverty %","Group"],
      [[r["code"],r["name"],r["gdp_pc_2021_usd"],r["gini"],r["poverty_pct"],r["group"]] for r in sorted(cov,key=lambda r:r["name"])],
      [7,19,16,8,10,12])

# ---- S5 wealth correlations ----
COV={r["code"]:r for r in cov}
cs=[c for c in nm if c in COV and all(COV[c].get(k,"") not in ("",None) for k in ("gdp_pc_2021_usd","gini","poverty_pct"))]
def col(key): return np.array([float(COV[c][key]) for c in cs])
gdp,gini,pov=col("gdp_pc_2021_usd"),col("gini"),col("poverty_pct")
s5=[]
for m in MODELS:
    ex=np.array([P(c,m,"2020-2025") for c in cs])
    row=[m]
    for x in (gdp,gini,pov): row+=[f"{pearsonr(ex,x)[0]:+.2f}",f"{spearmanr(ex,x)[0]:+.2f}"]
    s5.append(row)
sheet(wb,"S5 wealth correlations",["Baseline","GDP r","GDP ρ","Gini r","Gini ρ","Poverty r","Poverty ρ"],s5,[10]+[9]*6)

wb.save(os.path.join(DOCS,"PaperA_tables_twostep_v1.xlsx"))
print("wrote docs/PaperA_tables_twostep_v1.xlsx with sheets:",[ws.title for ws in wb.worksheets])

# ---------- (optional) emit the same tables into the manuscript (v6.0 -> v6.1) ----------
# This step is manuscript production, not reproduction: it inserts the tables into a specific
# Word file that is NOT part of this repository. It is skipped automatically when that file is
# absent, so the reproducible deliverable is the xlsx workbook written above.
_SRC_DOCX=os.path.join(DOCS,"Anchor_ML2b_PaperA_v6.0_twostep.docx")
if not os.path.exists(_SRC_DOCX):
    print("(manuscript docx not present — skipping the optional Word-table insertion step)")
    import sys as _sys; _sys.exit(0)
from docx import Document
CAP={"T1 methods comparison":"Table 1. Comparison with prior baseline-method studies.",
 "T2 slope table":"Table 2. Per-population pre-pandemic trend: intercept-slope at 2019 and slope-of-slopes (each with 95% CI), the shrinkage weight w, and the STTa and STTa+ shrunk slopes.",
 "T3 pooled excess":"Table 3. Pooled excess mortality across the 38 populations, by baseline and period (millions and P%).",
 "S1 age resolution":"Table S1. Age-resolution effect on the 2020–2025 excess (total excess deaths, millions), each model refitted at coarser age bands.",
 "S2 excess by location":"Table S2. Per-population 2020–2025 excess (P-score) by baseline, ordered by STTa+.",
 "S3 excess by period":"Table S3. Per-population excess (P-score) by period under STTa and STTa+.",
 "S4 vulnerability":"Table S4. Vulnerability classification and economic covariates.",
 "S5 wealth correlations":"Table S5. Correlation of the 2020–2025 excess with economic indicators, by baseline (Pearson r, Spearman ρ)."}
d=Document(os.path.join(DOCS,"Anchor_ML2b_PaperA_v6.0_twostep.docx"))
removed=0
for t in list(d.tables):
    if sum(1 for r in t.rows for c in r.cells if c.text.strip())==0:
        t._element.getparent().remove(t._element); removed+=1
d.add_page_break()
mainhdr=d.add_paragraph(); r=mainhdr.add_run("TABLES"); r.bold=True
for short,cap,header,rows in SPECS:
    if short.startswith("S1") :   # SI tables start here
        sih=d.add_paragraph(); rr=sih.add_run("SUPPLEMENTARY TABLES"); rr.bold=True
    p=d.add_paragraph(); rn=p.add_run(CAP[short]); rn.bold=True
    tb=d.add_table(rows=1,cols=len(header)); tb.style="Table Grid"
    for j,h in enumerate(header):
        c=tb.rows[0].cells[j]; c.text=str(h)
        for rr in c.paragraphs:
            for run in rr.runs: run.bold=True
    for row in rows:
        cells=tb.add_row().cells
        for j,v in enumerate(row): cells[j].text="" if v is None else str(v)
    d.add_paragraph()
OUT2=os.path.join(DOCS,"Anchor_ML2b_PaperA_v6.1_tables.docx"); d.save(OUT2)
d2=Document(OUT2)
print(f"removed {removed} empty shells; added {len(SPECS)} tables; wrote {OUT2} (now {len(d2.tables)} tables, {len(d2.paragraphs)} paragraphs)")
