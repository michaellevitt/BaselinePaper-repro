#!/usr/bin/env python3
"""
Per-country slope-of-slopes (SoS) and intercept-slope (the fitted moving-slope line value at 2019) with 95% CIs,
their population-weighted and unweighted means, and between-country HETEROGENEITY (Q, I^2, tau^2) — John's
point 9 (email 19 Jul 2026).  Unanchored parameters only (no '+'), consistent with the no-anchoring pivot.

Heterogeneity per parameter (DerSimonian-Laird): SE_i=(hi-lo)/(2*1.96); w_i=1/SE_i^2;
theta_fixed=Σw_iθ_i/Σw_i; Q=Σw_i(θ_i-theta_fixed)^2; df=k-1; I2=max(0,(Q-df)/Q);
tau2=max(0,(Q-df)/(Σw-Σw^2/Σw)); p from chi2(df).
Reads output/slope_of_slopes_CI.csv. Writes docs/Slope_of_slopes_heterogeneity_v1.xlsx + output/slope_heterogeneity.csv .
"""
import os, csv, numpy as np
from scipy import stats
HERE=os.path.dirname(os.path.abspath(__file__)); ROOT=os.path.dirname(HERE); OUTD=os.path.join(ROOT,"output"); DOCS=os.path.join(ROOT,"docs")
rows=[]; wmean={}
for r in csv.DictReader(open(os.path.join(OUTD,"slope_of_slopes_CI.csv"))):
    c=r["country"]
    if c.startswith("WEIGHTED"): wmean=dict(isl=float(r["u_islope2019"]),sos=float(r["u_sos"])); continue
    if c.startswith(("UNWEIGHTED","SD")) or not c: continue
    def g(k):
        try: return float(r[k])
        except: return np.nan
    rows.append(dict(c=c,name=r["name"],vuln=r["vulnerable"],pop=g("pop2019_millions"),
                     isl=g("u_islope2019"),isl_lo=g("u_islope2019_lo"),isl_hi=g("u_islope2019_hi"),
                     sos=g("u_sos"),sos_lo=g("u_sos_lo"),sos_hi=g("u_sos_hi")))
def het(theta,lo,hi):
    theta=np.array(theta); se=(np.array(hi)-np.array(lo))/(2*1.959964); w=1/se**2
    tf=np.sum(w*theta)/np.sum(w); Q=float(np.sum(w*(theta-tf)**2)); k=len(theta); df=k-1
    I2=max(0.0,(Q-df)/Q)*100 if Q>0 else 0.0
    C=np.sum(w)-np.sum(w**2)/np.sum(w); tau2=max(0.0,(Q-df)/C) if C>0 else 0.0
    p=float(stats.chi2.sf(Q,df)); return dict(Q=Q,df=df,I2=I2,tau2=tau2,p=p,tf=tf,k=k)
def popw(key):
    vv=[(r["pop"],r[key]) for r in rows if np.isfinite(r[key]) and np.isfinite(r["pop"])]
    W=sum(p for p,_ in vv); return sum(p*v for p,v in vv)/W
def uw(key): return float(np.nanmean([r[key] for r in rows if np.isfinite(r[key])]))
ISL=[r for r in rows if np.isfinite(r["isl"]) and np.isfinite(r["isl_lo"])]
SOS=[r for r in rows if np.isfinite(r["sos"]) and np.isfinite(r["sos_lo"])]
hISL=het([r["isl"] for r in ISL],[r["isl_lo"] for r in ISL],[r["isl_hi"] for r in ISL])
hSOS=het([r["sos"] for r in SOS],[r["sos_lo"] for r in SOS],[r["sos_hi"] for r in SOS])

# ---- console ----
print(f"n countries with intercept-slope CI: {hISL['k']} ; with SoS CI: {hSOS['k']}")
print(f"INTERCEPT-SLOPE @2019 (pct/yr):  pop-wtd {popw('isl'):.3f}  unwtd {uw('isl'):.3f}  | "
      f"Q={hISL['Q']:.1f} (df={hISL['df']}, p={hISL['p']:.1e})  I2={hISL['I2']:.0f}%  tau2={hISL['tau2']:.4f}")
print(f"SLOPE-OF-SLOPES (pct/yr^2):      pop-wtd {popw('sos'):.3f}  unwtd {uw('sos'):.3f}  | "
      f"Q={hSOS['Q']:.1f} (df={hSOS['df']}, p={hSOS['p']:.1e})  I2={hSOS['I2']:.0f}%  tau2={hSOS['tau2']:.4f}")

# ---- CSV for the paper ----
with open(os.path.join(OUTD,"slope_heterogeneity.csv"),"w",newline="") as f:
    w=csv.writer(f); w.writerow(["country","name","vulnerable","pop2019M","islope2019","islope_lo","islope_hi","sos","sos_lo","sos_hi"])
    for r in sorted(rows,key=lambda r:-(r["sos"] if np.isfinite(r["sos"]) else -9)):
        w.writerow([r["c"],r["name"],r["vuln"],r["pop"]]+[("" if not np.isfinite(r[k]) else round(r[k],3)) for k in
                    ["isl","isl_lo","isl_hi","sos","sos_lo","sos_hi"]])
    w.writerow([]); w.writerow(["POP-WEIGHTED MEAN","","","",round(popw("isl"),3),"","",round(popw("sos"),3),"",""])
    w.writerow(["UNWEIGHTED MEAN","","","",round(uw("isl"),3),"","",round(uw("sos"),3),"",""])
    w.writerow(["Q","","","",round(hISL["Q"],1),"","",round(hSOS["Q"],1),"",""])
    w.writerow(["I2 (%)","","","",round(hISL["I2"],0),"","",round(hSOS["I2"],0),"",""])
    w.writerow(["tau2","","","",round(hISL["tau2"],4),"","",round(hSOS["tau2"],4),"",""])
    w.writerow(["Q-test p","","","",f"{hISL['p']:.1e}","","",f"{hSOS['p']:.1e}","",""])

# ---- xlsx deliverable ----
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
wb=Workbook(); ws=wb.active; ws.title="Slope-of-slopes + heterogeneity"
HDR=PatternFill("solid",fgColor="1F4E79"); WHT=Font(color="FFFFFF",bold=True,size=9)
CEN=Alignment(horizontal="center"); LEFT=Alignment(horizontal="left")
thin=Side(style="thin",color="BFBFBF"); BORD=Border(left=thin,right=thin,top=thin,bottom=thin)
MOR=PatternFill("solid",fgColor="F8CBAD"); LES=PatternFill("solid",fgColor="C6E0B4"); SUM=PatternFill("solid",fgColor="D9E1F2")
ws["A1"]="Per-country pre-pandemic trend parameters (unanchored) and between-country heterogeneity — 38 populations"
ws["A1"].font=Font(bold=True,size=12)
ws["A2"]=("Intercept-slope = the fitted moving-slope line at 2019 (annual % change in the age-band death rate); "
          "slope-of-slopes = its trend (%/yr²; negative = decline accelerating). 95% CIs from the per-country regressions. "
          "Rows sorted by slope-of-slopes. Label fill: orange = more-, green = less-vulnerable.")
ws["A2"].font=Font(italic=True,size=9,color="555555"); ws.merge_cells("A2:G2"); ws.row_dimensions[2].height=40
hdr=["Population","2019 pop (M)","Intercept-slope\n2019 (%/yr)","95% CI","Slope-of-slopes\n(%/yr²)","95% CI","Vuln."]
for j,h in enumerate(hdr,1):
    c=ws.cell(4,j,h); c.fill=HDR; c.font=WHT; c.alignment=Alignment(horizontal="center",wrap_text=True,vertical="center"); c.border=BORD
ws.row_dimensions[4].height=30
r=5
for row in sorted(rows,key=lambda r:-(r["sos"] if np.isfinite(r["sos"]) else -9)):
    ci_i="" if not np.isfinite(row["isl_lo"]) else f"({row['isl_lo']:.2f}, {row['isl_hi']:.2f})"
    ci_s="" if not np.isfinite(row["sos_lo"]) else f"({row['sos_lo']:.3f}, {row['sos_hi']:.3f})"
    vals=[row["name"],f"{row['pop']:.1f}",("" if not np.isfinite(row['isl']) else f"{row['isl']:.2f}"),ci_i,
          ("" if not np.isfinite(row['sos']) else f"{row['sos']:.3f}"),ci_s,row["vuln"][:4]]
    for j,v in enumerate(vals,1):
        c=ws.cell(r,j,v); c.alignment=(LEFT if j==1 else CEN); c.border=BORD; c.font=Font(size=9)
    ws.cell(r,1).fill=(MOR if row["vuln"]=="more" else LES)
    r+=1
def srow(label,i_isl,i_sos,fill=SUM,bold=True):
    global r
    for j in range(1,8): ws.cell(r,j,"").fill=fill; ws.cell(r,j).border=BORD
    ws.cell(r,1,label).font=Font(bold=bold,size=9); ws.cell(r,1).alignment=LEFT
    ws.cell(r,3,i_isl).font=Font(bold=bold,size=9); ws.cell(r,3).alignment=CEN
    ws.cell(r,5,i_sos).font=Font(bold=bold,size=9); ws.cell(r,5).alignment=CEN
    r+=1
srow("Population-weighted mean",f"{popw('isl'):.3f}",f"{popw('sos'):.3f}")
srow("Unweighted mean",f"{uw('isl'):.3f}",f"{uw('sos'):.3f}")
srow(f"Q (Cochran)  [df={hISL['df']}]",f"{hISL['Q']:.1f}",f"{hSOS['Q']:.1f}")
srow("I²",f"{hISL['I2']:.0f}%",f"{hSOS['I2']:.0f}%")
srow("τ²",f"{hISL['tau2']:.4f}",f"{hSOS['tau2']:.4f}")
srow("Q-test p-value",f"{hISL['p']:.1e}",f"{hSOS['p']:.1e}")
ws.column_dimensions["A"].width=20
for col,wd in zip("BCDEFG",[11,14,15,14,15,7]): ws.column_dimensions[col].width=wd
ws.freeze_panes="A5"
OUT=os.path.join(DOCS,"Slope_of_slopes_heterogeneity_v1.xlsx"); wb.save(OUT)
print("\nwrote",OUT,"and output/slope_heterogeneity.csv")
